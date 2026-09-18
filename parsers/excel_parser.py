"""
Импорт расписания из Excel (.xlsx) в БД.

Запуск:
    # Только проверка, без записи:
    python parsers/excel_parser.py "1-4 курс 4 нед 21.09-26.09.xlsx"

    # Реальная запись (с бекапом и переименованием групп):
    python parsers/excel_parser.py "1-4 курс 4 нед 21.09-26.09.xlsx" --commit --rename-groups

    # Апдейт недели (удалить старые пары за эту неделю для этих групп и записать новые):
    python parsers/excel_parser.py "новый_файл.xlsx" --commit --rename-groups --reset-week
"""

import sys
import os
import re
import shutil
import sqlite3
from datetime import datetime

from openpyxl import load_workbook


BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, 'journal.db')
BACKUP_DIR = os.path.join(BASE_DIR, 'backups')
OUTPUT_DIR = os.path.join(BASE_DIR, 'parsers', 'output')


# ============================================================
#                 УТИЛИТЫ
# ============================================================

DAY_NAMES = {
    'понедельник': 1, 'вторник': 2, 'среда': 3, 'четверг': 4,
    'пятница': 5, 'суббота': 6, 'воскресенье': 7,
}


def cell_str(v):
    if v is None:
        return ''
    return str(v).strip()


def normalize_group_name(name):
    """
    'ИСП 21-9'       → 'ИСП 21'
    'ИСП 22-9, 24-11' → 'ИСП 22'
    'РУПО 11'        → 'РУПО 11'
    """
    first = name.split(',')[0].strip()
    first = re.sub(r'-\d+$', '', first)
    return first


def split_multi_teacher(raw):
    """
    'Лобанова Л.П. Тетенькин Д.А.' → 'Лобанова Л.П.'
    'Логов А.Г. Логов А.Г.' → 'Логов А.Г.'
    'Ватолина О.А.' → 'Ватолина О.А.'
    """
    if not raw:
        return raw
    s = raw.strip()
    pattern = r'[А-ЯЁ][а-яё\-]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.?'
    matches = re.findall(pattern, s)
    if matches:
        return matches[0]
    return s


def normalize_teacher_name(raw):
    """
    'Кусакина О.В,' → 'Кусакина О.В.'
    'Петров АВ'     → 'Петров А.В.'
    'Ватолина О.А.' → 'Ватолина О.А.'
    'Лобанова Л.П. Тетенькин Д.А.' → 'Лобанова Л.П.'
    """
    if not raw:
        return None
    s = split_multi_teacher(raw)
    s = s.strip()
    if s.endswith(','):
        s = s[:-1] + '.'
    m = re.match(r'^([А-ЯЁ][а-яё\-]+)\s+([А-ЯЁ])([А-ЯЁ])$', s)
    if m:
        return f"{m.group(1)} {m.group(2)}.{m.group(3)}."
    m = re.match(r'^([А-ЯЁ][а-яё\-]+)\s+([А-ЯЁ]\.)\s*([А-ЯЁ]\.?)$', s)
    if m:
        second = m.group(3)
        if not second.endswith('.'):
            second += '.'
        return f"{m.group(1)} {m.group(2)}{second}"
    return s


def clean_subject_name(raw):
    """
    Убирает мусорные префиксы/суффиксы из названия предмета.
    """
    if not raw:
        return raw
    s = raw.strip()

    # Убираем 'Вторник 22.09 8:30-10:00 1 ' в начале
    s = re.sub(r'^\S+\s+\d{2}\.\d{2}\s+\d{1,2}:\d{2}-\d{1,2}:\d{2}\s+\d+\s*', '', s)
    # Убираем '10:30-12:00 2 ' в начале
    s = re.sub(r'^\d{1,2}:\d{2}-\d{1,2}:\d{2}\s+\d+\s*', '', s)
    # Убираем одиночные '10:00-11:00 2', '14:30-15:30 6' целиком
    s = re.sub(r'^\d{1,2}:\d{2}-\d{1,2}:\d{2}\s+\d+$', '', s)
    # Убираем просто '2', '6', '7' (номер пары)
    if re.match(r'^\d+$', s):
        s = ''
    # Убираем хвост '\s+\d+\s+\d{1,2}:\d{2}-\d{1,2}:\d{2}\s+\d+$'
    s = re.sub(r'\s+\d+\s+\d{1,2}:\d{2}-\d{1,2}:\d{2}\s+\d+$', '', s)
    # Убираем хвост '\s+\d+$' (одиночный номер группы)
    s = re.sub(r'\s+\d+$', '', s)
    # Убираем префиксы-маркеры типа занятия
    low = s.lower()
    for prefix in ('д/з', 'д/ з', 'с/р', 'экзамен'):
        if low.startswith(prefix):
            s = s[len(prefix):].strip()
            break

    return s.strip()


def normalize_room(raw):
    """
    Возвращает (room_name, lesson_type_override).
    """
    if not raw:
        return None, None
    s = raw.strip()
    low = s.lower()

    if low in ('с/р', 'ср', 'с/ р'):
        return 'с/р', 'independent'
    if low in ('э', 'экзамен'):
        return 'э', 'exam'
    if low in ('д/з', 'дз') or low.startswith('д/з'):
        return 'Д/З', 'diff_credit'
    if low == 'библ':
        return 'Библиотека', None
    if low == 'акт зал':
        return 'акт зал', None
    if low == 'спорт зал':
        return 'Спорт зал', None

    if re.match(r'^\d+\.\d+$', s):
        return s.replace('.', ','), None

    return s, None


def detect_lesson_type_from_subject(subject, room_raw):
    """
    Определяет lesson_type.
    """
    if not subject:
        return 'lecture'
    low = subject.lower().strip()
    if low.startswith('экзамен'):
        return 'exam'
    if low.startswith('с/р'):
        return 'independent'
    if low.startswith('д/з') or low.startswith('д/ з'):
        return 'diff_credit'
    return 'lecture'


# ============================================================
#                 ЧТЕНИЕ ДАТ
# ============================================================

def read_dates(wb, year):
    if 'Даты' not in wb.sheetnames:
        raise RuntimeError("В файле нет листа 'Даты'")

    ws = wb['Даты']
    lookup = {}

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        day = None
        num = None
        for cell in row:
            v = cell_str(cell.value)
            if not v:
                continue
            low = v.lower()
            if low in DAY_NAMES:
                day = v
            elif re.match(r'^\d{2}\.\d{2}$', v):
                num = v

        if day and num:
            dd, mm = num.split('.')
            iso = f"{year}-{mm}-{dd}"
            lookup[f"{day} {num}"] = iso

    return lookup


# ============================================================
#                 ЧТЕНИЕ ЛИСТА
# ============================================================

def find_sections(ws):
    """
    Ищет в листе строки-шапки (где есть 'Дни', 'Часы', 'Ауд.').
    Правильно привязывает столбцы 'Ауд.' к ближайшей слева группе.
    """
    sections = []

    for row_idx in range(1, ws.max_row + 1):
        day_cols = []
        hour_cols = []
        room_cols = []

        for col_idx in range(1, ws.max_column + 1):
            v = cell_str(ws.cell(row=row_idx, column=col_idx).value)
            if v == 'Дни':
                day_cols.append(col_idx)
            elif v == 'Часы':
                hour_cols.append(col_idx)
            elif v == 'Ауд.':
                room_cols.append(col_idx)

        if not (day_cols and hour_cols and room_cols):
            continue

        def room_for(hour_col, next_day_col):
            upper = next_day_col if next_day_col else ws.max_column + 1
            for rc in room_cols:
                if hour_col < rc < upper:
                    return rc
            for rc in room_cols:
                if rc > hour_col:
                    return rc
            return None

        if len(day_cols) >= 2 and len(hour_cols) >= 2:
            left = {
                'day': day_cols[0], 'time': hour_cols[0],
                'room': room_for(hour_cols[0], day_cols[1]),
            }
            right = {
                'day': day_cols[1], 'time': hour_cols[1],
                'room': room_for(hour_cols[1], None),
            }
        else:
            left = {
                'day': day_cols[0], 'time': hour_cols[0],
                'room': room_for(hour_cols[0], None),
            }
            right = None

        def find_group_name(time_col, next_day_col):
            for r in (row_idx, row_idx + 1, row_idx + 2):
                for c in range(time_col + 1, next_day_col or ws.max_column + 1):
                    v = cell_str(ws.cell(row=r, column=c).value)
                    if not v:
                        continue
                    if v in ('Недели', 'Занятий', 'Дни', 'Часы', 'Ауд.', '№'):
                        continue
                    if re.match(r'^[А-ЯЁ]{2,5}\s+\d', v):
                        return v
            return None

        if right:
            left['name'] = find_group_name(hour_cols[0], day_cols[1])
            right['name'] = find_group_name(hour_cols[1], None)
            left['pair'] = hour_cols[0] + 2
            right['pair'] = hour_cols[1] + 2
            left['content_min'] = left['pair'] + 1
            left['content_max'] = (left['room'] - 1) if left['room'] else day_cols[1] - 1
            right['content_min'] = right['pair'] + 1
            right['content_max'] = (right['room'] - 1) if right['room'] else ws.max_column
        else:
            left['name'] = find_group_name(hour_cols[0], None)
            left['pair'] = hour_cols[0] + 2
            left['content_min'] = left['pair'] + 1
            left['content_max'] = (left['room'] - 1) if left['room'] else ws.max_column

        sections.append({
            'header_row': row_idx,
            'left': left,
            'right': right,
        })

    return sections


def parse_section(ws, section, dates_lookup, sheet_name):
    pairs = []

    for side_name in ('left', 'right'):
        side = section.get(side_name)
        if not side or not side.get('name'):
            continue

        group_name_raw = side['name']
        current_day = None
        row_idx = section['header_row'] + 1

        while row_idx <= ws.max_row:
            v = cell_str(ws.cell(row=row_idx, column=side['day']).value)
            if v == 'Дни':
                break

            date_cell = cell_str(ws.cell(row=row_idx, column=side['day']).value)
            if date_cell:
                m = re.match(r'^(\S+)\s+(\d{2}\.\d{2})$', date_cell)
                if m:
                    day_name = m.group(1)
                    day_num = DAY_NAMES.get(day_name.lower())
                    date_iso = dates_lookup.get(date_cell)
                    if date_iso and day_num:
                        current_day = {
                            'name': day_name,
                            'num': day_num,
                            'iso': date_iso,
                        }

            time_cell = cell_str(ws.cell(row=row_idx, column=side['time']).value)
            m_time = re.match(r'^(\d{1,2}:\d{2})-(\d{1,2}:\d{2})$', time_cell)

            if m_time and current_day:
                pair_num_raw = cell_str(ws.cell(row=row_idx, column=side['pair']).value)
                try:
                    pair_num = int(re.sub(r'\D', '', pair_num_raw))
                except (ValueError, TypeError):
                    pair_num = None

                if pair_num is not None:
                    subj = ''
                    teach = ''
                    for c in range(side['content_min'], side['content_max'] + 1):
                        v1 = cell_str(ws.cell(row=row_idx, column=c).value)
                        if v1:
                            subj = (subj + ' ' + v1).strip()
                    for c in range(side['content_min'], side['content_max'] + 1):
                        v2 = cell_str(ws.cell(row=row_idx + 1, column=c).value)
                        if v2:
                            teach = (teach + ' ' + v2).strip()

                    room_raw = ''
                    if side.get('room'):
                        room_raw = cell_str(ws.cell(row=row_idx, column=side['room']).value)

                    if not subj and not teach and not room_raw:
                        row_idx += 2
                        continue

                    pairs.append({
                        'sheet': sheet_name,
                        'group_name_raw': group_name_raw,
                        'day_name': current_day['name'],
                        'day_num': current_day['num'],
                        'date': current_day['iso'],
                        'pair_number': pair_num,
                        'time_start': m_time.group(1),
                        'time_end': m_time.group(2),
                        'subject_raw': subj,
                        'teacher_raw': teach,
                        'room_raw': room_raw,
                    })
                row_idx += 2
                continue

            row_idx += 1

    return pairs


# ============================================================
#                 БД
# ============================================================

def make_backup():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    dst = os.path.join(BACKUP_DIR, f'journal_{ts}_before_excel_import.db')
    shutil.copy2(DB_PATH, dst)
    return dst


def get_year_id(conn, year_name, start_date, end_date, rename_from=None):
    cur = conn.cursor()
    row = cur.execute(
        "SELECT id FROM academic_years WHERE name = ?", (year_name,)
    ).fetchone()
    if row:
        return row[0]

    if rename_from:
        row = cur.execute(
            "SELECT id FROM academic_years WHERE name = ?", (rename_from,)
        ).fetchone()
        if row:
            cur.execute(
                "UPDATE academic_years SET name = ?, start_date = ?, end_date = ?, is_current = 1 WHERE id = ?",
                (year_name, start_date, end_date, row[0])
            )
            return row[0]

    cur.execute(
        "INSERT INTO academic_years (name, start_date, end_date, is_current) VALUES (?, ?, ?, 1)",
        (year_name, start_date, end_date)
    )
    return cur.lastrowid


def match_or_create_group(conn, raw_name, rename_groups=False):
    cur = conn.cursor()
    row = cur.execute("SELECT id FROM groups WHERE name = ?", (raw_name,)).fetchone()
    if row:
        return row[0], False, False

    norm = normalize_group_name(raw_name)
    if norm != raw_name:
        row = cur.execute("SELECT id FROM groups WHERE name = ?", (norm,)).fetchone()
        if row:
            if rename_groups:
                cur.execute("UPDATE groups SET name = ? WHERE id = ?", (raw_name, row[0]))
                return row[0], False, True
            else:
                return row[0], False, False

    cur.execute("INSERT INTO groups (name) VALUES (?)", (raw_name,))
    return cur.lastrowid, True, False


def match_or_create_subject(conn, raw_name):
    cur = conn.cursor()
    name = raw_name.strip()
    if not name:
        return None, False
    row = cur.execute("SELECT id FROM subjects WHERE name = ?", (name,)).fetchone()
    if row:
        return row[0], False
    cur.execute("INSERT INTO subjects (name) VALUES (?)", (name,))
    return cur.lastrowid, True


def match_or_create_room(conn, raw_name):
    cur = conn.cursor()
    name = raw_name.strip()
    if not name:
        return None, False
    row = cur.execute("SELECT id FROM rooms WHERE name = ?", (name,)).fetchone()
    if row:
        return row[0], False
    cur.execute(
        "INSERT INTO rooms (name, building, capacity, room_type) VALUES (?, ?, ?, ?)",
        (name, 'учебный корпус', None, 'classroom')
    )
    return cur.lastrowid, True


def match_or_create_teacher(conn, raw_name):
    cur = conn.cursor()
    name = normalize_teacher_name(raw_name)
    if not name:
        return None, False
    row = cur.execute("SELECT id FROM teachers WHERE short_name = ?", (name,)).fetchone()
    if row:
        return row[0], False

    cur.execute(
        "INSERT INTO teachers (full_name, short_name, is_active) VALUES (?, ?, 1)",
        (name, name)
    )
    return cur.lastrowid, True


# ============================================================
#                 MAIN
# ============================================================

def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('xlsx', help='Путь к Excel-файлу расписания')
    parser.add_argument('--commit', action='store_true', help='Реально писать в БД')
    parser.add_argument('--rename-groups', action='store_true',
                        help='Переименовывать существующие группы')
    parser.add_argument('--reset-week', action='store_true',
                        help='Удалять существующие пары за неделю для групп из файла')
    parser.add_argument('--debug-sections', action='store_true',
                        help='Выводить координаты секций (диагностика)')
    args = parser.parse_args()

    if not os.path.exists(args.xlsx):
        print(f"❌ Файл не найден: {args.xlsx}")
        sys.exit(1)
    if not os.path.exists(DB_PATH):
        print(f"❌ БД не найдена: {DB_PATH}")
        sys.exit(1)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    report_file = os.path.join(OUTPUT_DIR, f'excel_parse_{ts}.txt')

    YEAR = 2026
    ACADEMIC_YEAR_NAME = '2026/2027'
    ACADEMIC_YEAR_START = '2026-09-01'
    ACADEMIC_YEAR_END = '2027-06-30'
    RENAME_FROM = '2025/2026'

    with open(report_file, 'w', encoding='utf-8') as out:
        def log(msg=''):
            print(msg)
            out.write(msg + '\n')

        log("=" * 80)
        log(f"EXCEL-ПАРСЕР РАСПИСАНИЯ")
        log(f"Файл:      {args.xlsx}")
        log(f"Режим:     {'COMMIT (запись в БД)' if args.commit else 'DRY-RUN'}")
        log(f"Год:       {ACADEMIC_YEAR_NAME}")
        log(f"Rename:    {args.rename_groups}")
        log(f"ResetWeek: {args.reset_week}")
        log(f"Дата:      {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        log("=" * 80)
        log()

        wb = load_workbook(args.xlsx, data_only=True)

        dates_lookup = read_dates(wb, YEAR)
        log(f"📅 Даты: {dates_lookup}")
        log()

        if args.debug_sections:
            log("=" * 80)
            log("DEBUG: Координаты секций")
            log("=" * 80)
            for sheet in wb.sheetnames:
                if sheet == 'Даты':
                    continue
                ws = wb[sheet]
                sections = find_sections(ws)
                log(f"\n[{sheet}] секций: {len(sections)}")
                for sec in sections:
                    log(f"  header_row={sec['header_row']}")
                    for side_name in ('left', 'right'):
                        side = sec.get(side_name)
                        if not side or not side.get('name'):
                            continue
                        log(f"    {side_name}: name='{side['name']}'")
                        log(f"      day={side['day']}, time={side['time']}, "
                            f"pair={side['pair']}, room={side['room']}")
                        log(f"      content=[{side['content_min']}..{side['content_max']}]")
            log()

        all_pairs = []
        for sheet in wb.sheetnames:
            if sheet == 'Даты':
                continue
            ws = wb[sheet]
            sections = find_sections(ws)
            log(f"📄 Лист '{sheet}': найдено секций — {len(sections)}")
            for sec in sections:
                sec_pairs = parse_section(ws, sec, dates_lookup, sheet)
                log(f"    • секция (строка {sec['header_row']}): "
                    f"left='{sec['left'].get('name')}', "
                    f"right='{sec['right'].get('name') if sec['right'] else '—'}', "
                    f"пар={len(sec_pairs)}")
                all_pairs.extend(sec_pairs)
            log()

        log(f"📊 Всего пар после отсева пустых: {len(all_pairs)}")
        log()

        day_count = {}
        for p in all_pairs:
            key = f"{p['day_name']} {p['date']}"
            day_count[key] = day_count.get(key, 0) + 1
        log("📅 Пар по дням:")
        for k in sorted(day_count.keys()):
            log(f"    {k:<25} : {day_count[k]}")
        log()

        conn = sqlite3.connect(DB_PATH)
        conn.execute("PRAGMA foreign_keys = ON")
        cur = conn.cursor()

        year_id = get_year_id(
            conn,
            ACADEMIC_YEAR_NAME,
            ACADEMIC_YEAR_START,
            ACADEMIC_YEAR_END,
            rename_from=RENAME_FROM if args.commit else None,
        )
        if args.commit:
            conn.commit()
        log(f"🎓 academic_year_id = {year_id} ({ACADEMIC_YEAR_NAME})")
        log()

        stats = {
            'groups_new': 0, 'groups_renamed': 0, 'groups_matched': 0,
            'subjects_new': 0, 'subjects_matched': 0,
            'rooms_new': 0, 'rooms_matched': 0,
            'teachers_new': 0, 'teachers_matched': 0, 'teachers_none': 0,
            'lessons_inserted': 0, 'lessons_deleted': 0,
        }

        log("=" * 80)
        log("ШАГ 1. Разбор пар и матчинг справочников")
        log("=" * 80)

        prepared = []

        for p in all_pairs:
            group_id, g_new, g_renamed = match_or_create_group(
                conn, p['group_name_raw'], rename_groups=args.rename_groups
            )
            if g_new:
                stats['groups_new'] += 1
            elif g_renamed:
                stats['groups_renamed'] += 1
            else:
                stats['groups_matched'] += 1

            room_name, type_override = normalize_room(p['room_raw'])
            lesson_type = type_override or detect_lesson_type_from_subject(
                p['subject_raw'], p['room_raw']
            )

            subject_id = None
            subject_name_clean = clean_subject_name(p['subject_raw'])
            if subject_name_clean:
                subject_id, s_new = match_or_create_subject(conn, subject_name_clean)
                if s_new:
                    stats['subjects_new'] += 1
                else:
                    stats['subjects_matched'] += 1

            teacher_id = None
            if p['teacher_raw']:
                teacher_id, t_new = match_or_create_teacher(conn, p['teacher_raw'])
                if t_new:
                    stats['teachers_new'] += 1
                else:
                    stats['teachers_matched'] += 1
            else:
                stats['teachers_none'] += 1

            room_id = None
            non_room_markers = ('с/р', 'э', 'Д/З')
            if room_name and room_name not in non_room_markers:
                room_id, r_new = match_or_create_room(conn, room_name)
                if r_new:
                    stats['rooms_new'] += 1
                else:
                    stats['rooms_matched'] += 1

            prepared.append({
                'pair': p,
                'group_id': group_id,
                'subject_id': subject_id,
                'subject_name_clean': subject_name_clean,
                'teacher_id': teacher_id,
                'room_id': room_id,
                'room_name_raw': room_name,
                'lesson_type': lesson_type,
            })

        if args.commit:
            conn.commit()
        else:
            conn.rollback()

        log()
        log("📊 Статистика матчинга:")
        for k, v in stats.items():
            log(f"    {k:<25} = {v}")
        log()

        log("=" * 80)
        log("ШАГ 1.5. Группы в файле")
        log("=" * 80)
        group_stats = {}
        for item in prepared:
            g = item['pair']['group_name_raw']
            group_stats[g] = group_stats.get(g, 0) + 1
        for g in sorted(group_stats.keys()):
            log(f"    {g:<40} : {group_stats[g]} пар")
        log()

        log("=" * 80)
        log("ШАГ 1.7. Что будет создано в БД")
        log("=" * 80)

        new_groups = set()
        new_subjects = set()
        new_teachers = set()
        new_rooms = set()

        for item in prepared:
            p = item['pair']
            _, g_new, _ = match_or_create_group(
                conn, p['group_name_raw'], rename_groups=args.rename_groups
            )
            if g_new:
                new_groups.add(p['group_name_raw'])

            if item.get('subject_name_clean'):
                _, s_new = match_or_create_subject(conn, item['subject_name_clean'])
                if s_new:
                    new_subjects.add(item['subject_name_clean'])

            if p['teacher_raw']:
                _, t_new = match_or_create_teacher(conn, p['teacher_raw'])
                if t_new:
                    new_teachers.add(normalize_teacher_name(p['teacher_raw']))

            if item['room_name_raw'] and item['room_name_raw'] not in ('с/р', 'э', 'Д/З'):
                _, r_new = match_or_create_room(conn, item['room_name_raw'])
                if r_new:
                    new_rooms.add(item['room_name_raw'])

        conn.rollback()

        log(f"\n🆕 Новые группы ({len(new_groups)}):")
        for g in sorted(new_groups):
            log(f"    • {g}")

        log(f"\n🆕 Новые предметы ({len(new_subjects)}, первые 40):")
        for s in sorted(new_subjects)[:40]:
            log(f"    • {s}")
        if len(new_subjects) > 40:
            log(f"    ... и ещё {len(new_subjects) - 40}")

        log(f"\n🆕 Новые преподаватели ({len(new_teachers)}):")
        for t in sorted(new_teachers):
            log(f"    • {t}")

        log(f"\n🆕 Новые аудитории ({len(new_rooms)}):")
        for r in sorted(new_rooms):
            log(f"    • {r}")

        log(f"\n🔤 Переименование групп (будет при --rename-groups):")
        for raw_name in sorted(set(p['group_name_raw'] for p in all_pairs)):
            norm = normalize_group_name(raw_name)
            if norm != raw_name:
                row = conn.execute(
                    "SELECT id, name FROM groups WHERE name = ?", (norm,)
                ).fetchone()
                if row:
                    row_id = row[0]
                    row_name = row[1]
                    log(f"    • '{row_name}' (id={row_id}) → '{raw_name}'")
        log()

        log("=" * 80)
        log("ШАГ 2. Примеры разбора пар (первые 20)")
        log("=" * 80)
        for item in prepared[:20]:
            p = item['pair']
            log(f"  [{p['sheet']}] {p['group_name_raw']} | {p['date']} ({p['day_name']}) "
                f"пара {p['pair_number']} ({p['time_start']}-{p['time_end']})")
            log(f"      предмет: '{p['subject_raw']}' → '{item['subject_name_clean']}'")
            log(f"      препод:  '{p['teacher_raw']}' → {normalize_teacher_name(p['teacher_raw'])}")
            log(f"      ауд:     '{p['room_raw']}' → {item['room_name_raw']}")
            log(f"      тип:     {item['lesson_type']}")
            log()

        if args.commit:
            log("=" * 80)
            log("ШАГ 3. Запись в БД")
            log("=" * 80)

            backup = make_backup()
            log(f"💾 Бекап: {backup}")
            log()

            if args.reset_week:
                dates = sorted(set(item['pair']['date'] for item in prepared))
                groups = sorted(set(item['group_id'] for item in prepared))
                if dates and groups:
                    d_min, d_max = dates[0], dates[-1]
                    placeholders_g = ','.join('?' for _ in groups)
                    q = (f"DELETE FROM schedule_lessons "
                         f"WHERE date BETWEEN ? AND ? "
                         f"AND group_id IN ({placeholders_g})")
                    cur.execute(q, [d_min, d_max] + groups)
                    stats['lessons_deleted'] = cur.rowcount
                    log(f"🗑️  Удалено старых пар: {stats['lessons_deleted']} "
                        f"(диапазон {d_min}…{d_max}, группы {groups})")
                    log()

            for item in prepared:
                p = item['pair']
                cur.execute('''
                    INSERT INTO schedule_lessons
                    (academic_year_id, date, day_of_week, pair_number,
                     time_start, time_end,
                     group_id, subject_id,
                     teacher_id, teacher_name_raw,
                     room_id, room_name_raw,
                     lesson_type, original_group_name, source_file)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    year_id, p['date'], p['day_num'], p['pair_number'],
                    p['time_start'], p['time_end'],
                    item['group_id'], item['subject_id'],
                    item['teacher_id'],
                    normalize_teacher_name(p['teacher_raw']) or p['teacher_raw'],
                    item['room_id'], item['room_name_raw'],
                    item['lesson_type'], p['group_name_raw'],
                    os.path.basename(args.xlsx),
                ))
                stats['lessons_inserted'] += 1

            conn.commit()
            log(f"✅ Записано пар: {stats['lessons_inserted']}")
            log()

        log("=" * 80)
        log("ИТОГО")
        log("=" * 80)
        for k, v in stats.items():
            log(f"  {k:<25} = {v}")
        log()
        log(f"📁 Отчёт: {report_file}")
        log("=" * 80)

        conn.close()

    print(f"\n📁 Отчёт: {report_file}")
    if not args.commit:
        print("⚠️  Это был DRY-RUN. Ничего не записано.")
        print("   Для записи добавь: --commit --rename-groups")


if __name__ == '__main__':
    main()